"""연도별 문체부 개최계획의 병합 머리글·단위·일정을 정규화하고 품질을 기록한다."""

import json
import re
from datetime import date
from pathlib import Path

import polars as pl
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from crowdcast.data import FESTIVAL_DTYPES, FESTIVAL_SCHEMA, SIDO_ALIASES, TYPE_MAPPING, quality_report

# 셀 값 해석(문구·단위·일수·일정·개최 달)은 mcst_values에 있다
from crowdcast.data.mcst_values import (
    dates,
    duration,
    integer,
    key,
    planned_month,
    text,
)
from crowdcast.paths import DATA, PROCESSED

FIELDS = {
    "sido": r"^(시도명|광역.*단체명)$",
    "sigungu_name": r"^(시군구명|기초.*단체명)$",
    "festival_name": r"^(축제명|행사명|축제명칭)$",
    "type_raw": r"축제(유형|종류)",
    "venue": r"개최장소|장소명",
    "host": r"주최/주관|전담조직명|^주최$",
    "period": r"개최기간",
    "start": r"시작일",
    "end": r"종료일",
    "days": r"총일수|^일수$",
    "period_note": r"개최기간.*비고",
}


# 시트명과 행 값에 같은 시도 정규화를 적용해 주석·다른 지역 행을 구별한다.
def sido_name(value: object) -> str | None:
    name = key(value).replace("강원도", "강원특별자치도").replace("전라북도", "전북특별자치도")
    name = name.replace("제주도", "제주특별자치도").replace("세종특별자치시청", "세종특별자치시")
    return SIDO_ALIASES.get(name, text(name))


# 실제 병합 범위만 복원해 빈 지역 셀을 앞 행사 값으로 잘못 채우지 않는다.
def sheet_rows(sheet: Worksheet) -> tuple[list[list[object]], dict[tuple[int, int], tuple[int, int]]]:
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    merged = {}
    for area in sheet.merged_cells.ranges:
        for r in range(area.min_row - 1, area.max_row):
            for c in range(area.min_col - 1, area.max_col):
                rows[r][c] = rows[area.min_row - 1][area.min_col - 1]
                merged[r, c] = (area.min_row - 1, area.min_col - 1)
    return rows, merged


# 행사명 머리글을 기준으로 병합된 여러 행의 열 제목을 결합한다.
def headers(rows: list[list[object]], merged: dict) -> tuple[int, list[str]] | None:
    for r, row in enumerate(rows[:40]):
        for c, value in enumerate(row):
            if not re.fullmatch(FIELDS["festival_name"], key(value)):
                continue
            last = max((rr for (rr, cc), origin in merged.items() if origin == (r, c)), default=r)
            # 행사명이 비어 있는 바로 아래 단위 행도 병합 머리글의 일부로 읽는다.
            if (
                last + 1 < len(rows)
                and rows[last + 1][c] is None
                and any(re.search(r"단위\s*[:：]|백만원", str(v)) for v in rows[last + 1] if v is not None)
            ):
                last += 1
            titles = []
            for col in range(len(row)):
                values = [str(rows[rr][col]) for rr in range(r, last + 1) if rows[rr][col] is not None]
                if r and re.search(r"방문|예산|단위", str(rows[r - 1][col])):
                    values.insert(0, str(rows[r - 1][col]))
                titles.append(" / ".join(dict.fromkeys(values)))
            return last, titles
    return None


# 전담 조직명에서 연락처·담당자 기재 부분을 버리고 조직 부분만 보존한다.
def host_name(value: object) -> str | None:
    contact = r"담당|연락|위원장|회장|대표자|주무관|팀장|과장|[\w.+-]+@[\w.-]+|\d{2,4}[-) ]\d{3,4}[- ]\d{4}"
    raw = re.split(contact, str(value or ""))[0]
    raw = re.sub(r"\([^)]*$", "", raw)
    raw = re.sub(r"((?:위원회|협의회|조직위|재단|시|군|구|청|과|팀))\s+[가-힣]{2,4}\s*$", r"\1", raw)
    return text(raw.strip(" /,;\n"))


# 열 위치 대신 제목으로 값을 찾고 원본의 행 번호를 함께 반환한다.
def parse_sheet(sheet: Worksheet, year: int, source_file: str) -> list[dict]:
    rows, merged = sheet_rows(sheet)
    layout = headers(rows, merged)
    if layout is None:
        return []
    last, titles = layout
    columns = {
        name: next((i for i, title in enumerate(titles) if re.search(pattern, key(title))), None)
        for name, pattern in FIELDS.items()
    }
    # 합계 열을 우선하며 부분 집계로만 구성된 양식은 두 원문 제목을 모두 남긴다.
    visitors = [i for i, title in enumerate(titles) if re.search(r"방문객|방문자", title)]
    totals = [i for i in visitors if re.search(r"합계|전체", key(titles[i]))]
    visitors = totals[:1] or [i for i in visitors if re.search(r"내국인|외국인", titles[i])] or visitors[:1]
    budget = next(
        (
            i
            for i, title in enumerate(titles)
            if "예산" in title and not re.search(r"국비|지방비|기타|시.?도비|구.?군비", title)
        ),
        None,
    )
    if budget is None:
        budget = next(
            (
                i - 1
                for i, title in enumerate(titles)
                if i and "예산" in title and key(titles[i - 1]) == "합계"
            ),
            None,
        )
    budget_heading = " ".join(t for t in titles if "예산" in t or "백만원" in t)
    shared_counts = {
        origin for (rr, cc), origin in merged.items() if rr > origin[0] > last and cc in visitors
    }
    # 시도별 파일의 빈 시도는 명확히 알려진 시트명에서만 보완한다.
    sheet_sido = sido_name(sheet.title)
    if sheet_sido not in SIDO_ALIASES.values():
        sheet_sido = None
    records = []
    for r in range(last + 1, len(rows)):
        row = rows[r]
        values = {name: row[col] if col is not None else None for name, col in columns.items()}
        name = text(values["festival_name"])
        if not name or name in {"축제명", "합계", "총계", "소계", "<보기>"}:
            continue
        name_cell = (r, columns["festival_name"])
        if merged.get(name_cell, name_cell) != name_cell:
            continue
        sido = sido_name(values["sido"]) or sheet_sido
        if sheet_sido and sido != sheet_sido:
            continue
        if not text(values["sido"]) and not text(values["period"]) and not sheet_sido:
            continue
        # 분리 연월일은 실제 일자가 있을 때만 조합하고 일수 불일치·반복 일정은 보류한다.
        period, col, stated_days = values["period"], columns["period"], values["days"]
        if col is not None and col + 1 < len(titles) and titles[col] == titles[col + 1]:
            if merged.get((r, col + 1)) != (r, col):
                stated_days = row[col + 1]
        date_text = str(period) if period is not None else None
        start, end, days = dates(period, year, stated_days)
        month = planned_month(period, year)
        if columns["start"] is not None and columns["end"] is not None:
            ends, months, original = [], [], {}
            for field in ("start", "end"):
                col = columns[field]
                parts = row[col : col + 3] if titles[col].endswith("년") else [row[col]]
                original[field] = parts
                months.append(integer(parts[1]) if len(parts) == 3 else planned_month(parts[0], year))
                try:
                    ends.append(
                        date(*(int(v) for v in parts)) if len(parts) == 3 else dates(parts[0], year)[0]
                    )
                except (ValueError, TypeError):
                    ends.append(None)
            start, end = ends
            start, end = (None, None) if start and end and end < start else (start, end)
            original.update(days=values["days"], note=values["period_note"])
            date_text = json.dumps(original, ensure_ascii=False, default=str)
            month = months[0] if months[0] == months[1] and months[0] in range(1, 13) else None
            stated_days = duration(values["days"])
            days = (end - start).days + 1 if start and end else stated_days
            if (stated_days and days != stated_days) or re.search(
                r"미정|취소|주\s*1회|월\s*1회|매주|매월|기간\s*중", str(values["period_note"])
            ):
                start, end, days = None, None, stated_days
        # 가로 병합된 방문객 합계는 중복 가산하지 않고 지역·유형의 원문은 보존한다.
        origins = list(dict.fromkeys(merged.get((r, col), (r, col)) for col in visitors))
        counts = [
            integer(rows[rr][cc], titles[cc]) if (rr, cc) not in shared_counts else None for rr, cc in origins
        ]
        records.append(
            {
                "year": year,
                "sido": sido,
                "sigungu_name": text(values["sigungu_name"]),
                "festival_name": name,
                "type_raw": text(values["type_raw"]),
                "type": TYPE_MAPPING.get(key(values["type_raw"]), "기타"),
                "start_date": start,
                "end_date": end,
                "days": days if days and days >= 1 else None,
                "date_text": date_text,
                "planned_month": month,
                "venue": text(values["venue"]),
                "host": host_name(values["host"]),
                "budget_krw": (
                    integer(row[budget], budget_heading)
                    if budget is not None and "원" in budget_heading + str(row[budget])
                    else None
                ),
                "visitors_announced": sum(counts) if counts and all(n is not None for n in counts) else None,
                "visitors_announced_meaning": " + ".join(titles[col] for col in visitors) or None,
                "source_file": source_file,
                "source_sheet": sheet.title,
                "source_row": r + 1,
            }
        )
    return records


# 원본은 읽기만 하고 총괄·보기 시트는 행사 머리글 유무로 제외한다.
def parse_workbook(path: Path, year: int, source_file: str | None = None) -> list[dict]:
    workbook = load_workbook(path, data_only=True)
    try:
        return [record for sheet in workbook for record in parse_sheet(sheet, year, source_file or path.name)]
    finally:
        workbook.close()


# 고정 경로에 검증된 행사 표와 품질 보고서를 저장하고 연도별 최소 행 수를 확인한다.
def main() -> None:
    files = sorted(path for path in DATA.glob("20*_festival/**/*.xlsx") if not path.name.startswith("~$"))
    if not files:
        raise FileNotFoundError(f"개최계획 xlsx 없음: {DATA}")
    records = [
        row
        for path in files
        for row in parse_workbook(
            path, int(path.relative_to(DATA).parts[0][:4]), path.relative_to(DATA).as_posix()
        )
    ]
    frame = FESTIVAL_SCHEMA.validate(pl.DataFrame(records, schema=FESTIVAL_DTYPES), lazy=True)
    PROCESSED.mkdir(parents=True, exist_ok=True)
    (PROCESSED / "mcst_festivals_qc.md").write_text(quality_report(frame), encoding="utf-8")
    counts = dict(frame.group_by("year").len().iter_rows())
    if any(counts.get(year, 0) < 300 for year in range(2017, 2027)):
        raise ValueError("연도별 300행 기준 미달: mcst_festivals_qc.md 확인")
    frame.write_parquet(PROCESSED / "mcst_festivals.parquet")
    print(f"문체부 행사 {frame.height}행 저장: {PROCESSED}")


if __name__ == "__main__":
    main()
